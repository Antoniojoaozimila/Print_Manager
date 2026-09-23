#!/usr/bin/env python3
"""Preenche o modelo original da Matriz de Controlo Diário TI, preservando
gráficos, cores, grelhas, fórmulas, validações e todas as abas."""
from __future__ import annotations

import argparse
import json
import re
import sys
from copy import copy
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

DEFAULT_TECHS = ["Zimila", "Elton", "Edna"]
ASSIST_MAX = 1001
PROJECTO_MAX = 501
AQUISICAO_MAX = 501


def parse_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", text)
    if m:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", text)
    if m:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return None


def parse_time(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    text = str(value).strip()
    if not text:
        return None
    iso = re.search(r"T(\d{2}:\d{2}(?::\d{2})?)", text)
    if iso:
        text = iso.group(1)
    m = re.match(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?", text)
    if not m:
        return None
    return time(int(m.group(1)), int(m.group(2)), int(m.group(3) or 0))


def copy_style(src, dest):
    dest.font = copy(src.font)
    dest.fill = copy(src.fill)
    dest.border = copy(src.border)
    dest.alignment = copy(src.alignment)
    dest.number_format = src.number_format
    dest.protection = copy(src.protection)


def clone_row_style(ws: Worksheet, src_row: int, dest_row: int, max_col: int):
    src_dim = ws.row_dimensions[src_row]
    dest_dim = ws.row_dimensions[dest_row]
    dest_dim.height = src_dim.height
    for col in range(1, max_col + 1):
        copy_style(ws.cell(src_row, col), ws.cell(dest_row, col))


def set_a_formula(ws: Worksheet, row: int):
    ws.cell(row, 1).value = f'=IF(B{row}="","",ROW()-1)'


def set_duracao_formula(ws: Worksheet, row: int):
    ws.cell(row, 6).value = f'=IF(OR(D{row}="",E{row}=""),"",ROUND(MOD(E{row}-D{row},1)*1440,0))'


def set_doc_formula(ws: Worksheet, row: int):
    ws.cell(row, 20).value = (
        f'=IF(B{row}="","",IF(AND(I{row}<>"",O{row}<>"",P{row}<>"",Q{row}<>"",R{row}<>"",S{row}<>""),"Completa","Pendente"))'
    )


def ensure_table_row(ws: Worksheet, row: int, template_row: int, max_col: int, setup):
    if row <= template_row:
        return
    clone_row_style(ws, template_row, row, max_col)
    setup(ws, row)


def clear_assistencia_row(ws: Worksheet, row: int):
    for col in range(1, 19):
        if col == 6:
            continue
        ws.cell(row, col).value = None
    set_a_formula(ws, row)
    if not ws.cell(row, 6).value:
        set_duracao_formula(ws, row)


def clear_projecto_row(ws: Worksheet, row: int):
    for col in range(2, 16):
        ws.cell(row, col).value = None
    set_a_formula(ws, row)


def write_assistencias(ws: Worksheet, rows: list):
    last_template = ASSIST_MAX
    for r in range(2, last_template + 1):
        clear_assistencia_row(ws, r)

    for i, item in enumerate(rows):
        row = 2 + i
        ensure_table_row(ws, row, last_template, 18, set_duracao_formula)
        if row > last_template:
            set_a_formula(ws, row)
            set_duracao_formula(ws, row)

        ws.cell(row, 1).value = item.get("numero")
        data = parse_date(item.get("data"))
        ws.cell(row, 2).value = data
        ws.cell(row, 2).number_format = "dd/mm/yyyy"
        ws.cell(row, 3).value = item.get("tecnico") or None
        inicio = parse_time(item.get("hora_inicio"))
        fim = parse_time(item.get("hora_fim"))
        ws.cell(row, 4).value = inicio
        ws.cell(row, 4).number_format = "hh:mm"
        ws.cell(row, 5).value = fim
        ws.cell(row, 5).number_format = "hh:mm"
        set_duracao_formula(ws, row)
        ws.cell(row, 7).value = item.get("departamento") or None
        ws.cell(row, 8).value = item.get("provincia") or None
        ws.cell(row, 9).value = item.get("colaborador_assistido") or None
        ws.cell(row, 10).value = item.get("tipo_solicitacao") or None
        ws.cell(row, 11).value = item.get("problema") or None
        ws.cell(row, 12).value = item.get("resolucao") or None
        ws.cell(row, 13).value = item.get("estado") or None
        urg = item.get("urgencia")
        if urg in (True, 1, "1", "Sim", "sim", "SIM"):
            ws.cell(row, 14).value = "Sim"
        elif urg in (False, 0, "0", "Não", "Nao", "não", "nao"):
            ws.cell(row, 14).value = "Não"
        else:
            ws.cell(row, 14).value = urg or "Não"
        ws.cell(row, 15).value = item.get("descricao_urgencia") or None
        chamadas = item.get("num_chamadas")
        ws.cell(row, 16).value = int(chamadas) if chamadas not in (None, "") else 1
        ws.cell(row, 17).value = item.get("meio_solicitacao") or None
        ws.cell(row, 18).value = item.get("observacoes") or None

    used = 2 + len(rows) - 1 if rows else 1
    ws.auto_filter.ref = f"A1:R{max(last_template, used)}"
    ws.freeze_panes = "A2"


def write_projectos(ws: Worksheet, rows: list):
    last_template = PROJECTO_MAX
    for r in range(2, last_template + 1):
        clear_projecto_row(ws, r)

    for i, item in enumerate(rows):
        row = 2 + i
        ensure_table_row(ws, row, last_template, 15, set_a_formula)
        if row > last_template:
            set_a_formula(ws, row)

        ws.cell(row, 1).value = item.get("numero")
        ws.cell(row, 2).value = parse_date(item.get("data"))
        ws.cell(row, 2).number_format = "dd/mm/yyyy"
        ws.cell(row, 3).value = item.get("responsavel") or None
        ws.cell(row, 4).value = item.get("projecto_sistema") or None
        ws.cell(row, 5).value = item.get("tarefa") or None
        ws.cell(row, 6).value = parse_date(item.get("data_atribuicao"))
        ws.cell(row, 6).number_format = "dd/mm/yyyy"
        ws.cell(row, 7).value = parse_date(item.get("prazo"))
        ws.cell(row, 7).number_format = "dd/mm/yyyy"
        ws.cell(row, 8).value = item.get("fase_actual") or None
        pct = item.get("percentagem_conclusao")
        if pct in (None, ""):
            ws.cell(row, 9).value = 0
        else:
            ws.cell(row, 9).value = max(0, min(100, float(pct))) / 100.0
        ws.cell(row, 9).number_format = "0%"
        ws.cell(row, 10).value = item.get("alteracoes_solicitadas") or None
        ws.cell(row, 11).value = parse_date(item.get("data_alteracao"))
        ws.cell(row, 11).number_format = "dd/mm/yyyy"
        ws.cell(row, 12).value = item.get("descricao_alteracao") or None
        ws.cell(row, 13).value = item.get("accao_realizada") or None
        ws.cell(row, 14).value = item.get("estado") or None
        ws.cell(row, 15).value = item.get("observacoes") or None

    used = 2 + len(rows) - 1 if rows else 1
    ws.auto_filter.ref = f"A1:O{max(last_template, used)}"
    ws.freeze_panes = "A2"


def extra_tecnicos(assistencias, projectos):
    seen = []
    known = set(DEFAULT_TECHS)
    for item in assistencias:
        name = (item.get("tecnico") or "").strip()
        if name and name not in known and name not in seen:
            seen.append(name)
    for item in projectos:
        name = (item.get("responsavel") or "").strip()
        if name and name not in known and name not in seen:
            seen.append(name)
    return seen


def tech_formulas(row: int) -> dict:
    assist = "'ASSISTÊNCIA DIÁRIA'"
    proj = "'PROJECTOS E TAREFAS'"
    return {
        2: f"=COUNTIF({assist}!$C$2:$C$1001,A{row})",
        3: f'=COUNTIFS({assist}!$C$2:$C$1001,A{row},{assist}!$M$2:$M$1001,"Concluído")',
        4: (
            f'=COUNTIFS({assist}!$C$2:$C$1001,A{row},{assist}!$M$2:$M$1001,"Pendente")'
            f'+COUNTIFS({assist}!$C$2:$C$1001,A{row},{assist}!$M$2:$M$1001,"Em Progresso")'
        ),
        5: f'=COUNTIFS({assist}!$C$2:$C$1001,A{row},{assist}!$N$2:$N$1001,"Sim")',
        6: f"=SUMIF({assist}!$C$2:$C$1001,A{row},{assist}!$F$2:$F$1001)",
        7: f"=IFERROR(F{row}/B{row},0)",
        8: f"=COUNTIF({proj}!$C$2:$C$501,A{row})",
    }


def add_dashboard_tecnicos(ws: Worksheet, extras: list):
    last = 7
    if extras:
        extra_slots = 2
        overflow = max(0, len(extras) - extra_slots)
        if overflow:
            ws.insert_rows(10, overflow)
        for i, name in enumerate(extras):
            dest = 8 + i
            clone_row_style(ws, 7, dest, 8)
            ws.cell(dest, 1).value = name
            for col, formula in tech_formulas(dest).items():
                ws.cell(dest, col).value = formula
        last = 7 + len(extras)

    if ws._charts:
        chart = ws._charts[0]
        if chart.series:
            series = chart.series[0]
            if series.val is not None and series.val.numRef is not None:
                series.val.numRef.f = f"'DASHBOARD IT'!$B$5:$B${last}"
            if series.cat is not None and series.cat.strRef is not None:
                series.cat.strRef.f = f"'DASHBOARD IT'!$A$5:$A${last}"


def expand_listas_tecnicos(wb, extras: list):
    if not extras:
        return
    ws = wb["LISTAS"]
    for i, name in enumerate(extras):
        dest = 5 + i
        copy_style(ws.cell(4, 1), ws.cell(dest, 1))
        ws.cell(dest, 1).value = name
    last = 4 + len(extras)
    new_ref = f"LISTAS!$A$2:$A${last}"
    new_ref_eq = f"={new_ref}"
    old_refs = {"LISTAS!$A$2:$A$4", "=LISTAS!$A$2:$A$4"}
    for sheet_name in ("ASSISTÊNCIA DIÁRIA", "PROJECTOS E TAREFAS", "AQUISIÇÕES"):
        sh = wb[sheet_name]
        for dv in sh.data_validations.dataValidation:
            if (dv.formula1 or "") in old_refs:
                dv.formula1 = new_ref_eq if str(dv.formula1).startswith("=") else new_ref


def ensure_aquisicoes_formulas(ws: Worksheet):
    for row in range(2, AQUISICAO_MAX + 1):
        if not ws.cell(row, 1).value:
            set_a_formula(ws, row)
        if not ws.cell(row, 20).value:
            set_doc_formula(ws, row)
    ws.auto_filter.ref = f"A1:U{AQUISICAO_MAX}"
    ws.freeze_panes = "A2"


def exportar(template: Path, data_path: Path, output: Path):
    payload = json.loads(data_path.read_text(encoding="utf-8"))
    assistencias = payload.get("assistencias") or []
    projectos = payload.get("projectos") or []
    extras = extra_tecnicos(assistencias, projectos)

    wb = load_workbook(template)
    write_assistencias(wb["ASSISTÊNCIA DIÁRIA"], assistencias)
    write_projectos(wb["PROJECTOS E TAREFAS"], projectos)
    ensure_aquisicoes_formulas(wb["AQUISIÇÕES"])
    expand_listas_tecnicos(wb, extras)
    add_dashboard_tecnicos(wb["DASHBOARD IT"], extras)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def main(argv=None):
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--data", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args(argv)
    try:
        exportar(Path(args.template), Path(args.data), Path(args.output))
    except Exception as exc:  # noqa: BLE001
        print(str(exc), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
